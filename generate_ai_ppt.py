from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "The Current State of AI & Its Impact on Daily Life"
    subtitle.text = "An Overview of Artificial Intelligence Today and Tomorrow\n\nGenerated for you by Antigravity"

    slides_data = [
        ("What is the 'Current State' of AI?", 
         "• Generative AI: Creating text, images, and video (e.g., ChatGPT, Midjourney).\n"
         "• Large Language Models (LLMs): Understanding and generating human-like text.\n"
         "• Computer Vision: Analyzing and understanding the visual world.\n"
         "• Shift from narrow, task-specific AI to broader, more generalized capabilities."),
        ("The AI Boom: Why Now?",
         "• Massive Data: Abundance of digital text, images, and interactions.\n"
         "• Compute Power: Advanced GPUs capable of massive parallel processing.\n"
         "• Algorithmic Breakthroughs: The 'Transformer' architecture revolutionizing machine learning.\n"
         "• Heavy Investment: Billions of dollars pouring into AI research and startups."),
        ("AI in the Workplace",
         "• Automation of Routine Tasks: Data entry, scheduling, and basic customer service.\n"
         "• AI 'Copilots': Assisting in coding, writing, and design.\n"
         "• Enhanced Decision Making: Data analytics and predictive modeling.\n"
         "• Shift in Skills: Growing demand for AI literacy and prompt engineering."),
        ("AI in Healthcare",
         "• Diagnostics: Assisting doctors in reading X-rays and MRIs with high accuracy.\n"
         "• Drug Discovery: Accelerating the identification of new compounds and treatments.\n"
         "• Personalized Medicine: Tailoring treatments based on genetic and patient data.\n"
         "• Virtual Health Assistants: Providing 24/7 patient support and triage."),
        ("AI in Education",
         "• Personalized Learning: Adapting to individual student's pace and style.\n"
         "• Intelligent Tutors: Providing instant feedback and explanations.\n"
         "• Administrative Automation: Grading and lesson planning assistance for teachers.\n"
         "• Accessibility: Real-time translation and transcription for diverse needs."),
        ("AI in Daily Life",
         "• Smart Homes: Thermostats, lighting, and security adapting to your habits.\n"
         "• Virtual Assistants: Siri, Alexa, and Google Assistant becoming more conversational.\n"
         "• Recommendation Algorithms: Curating content on Netflix, Spotify, and social media.\n"
         "• Transportation: Optimized routing, autonomous driving features, and smart traffic management."),
        ("Creative Arts & Entertainment",
         "• Visual Arts: AI-generated images and concepts for marketing and design.\n"
         "• Writing: Brainstorming, drafting, and editing text.\n"
         "• Music and Video: Generating background music, voiceovers, and even short video clips.\n"
         "• Gaming: Dynamic NPCs, procedural content generation, and adaptive difficulty."),
        ("Ethical & Social Implications",
         "• Bias and Fairness: AI inheriting prejudices present in training data.\n"
         "• Job Displacement vs. Creation: The transition of the workforce.\n"
         "• Misinformation: Deepfakes, automated propaganda, and loss of trust.\n"
         "• Privacy Concerns: Massive data collection required to fuel AI models."),
        ("The Future of AI",
         "• Toward AGI (Artificial General Intelligence): Systems that equal or exceed human intelligence across a wide range of tasks.\n"
         "• Robotics Integration: AI powering physical systems for manufacturing and home assistance.\n"
         "• Regulation & Governance: Developing laws to ensure safe and responsible AI development.\n"
         "• Human-AI Collaboration: A future where AI augments rather than replaces human capabilities."),
        ("Conclusion",
         "• AI is no longer science fiction; it is a fundamental part of our current reality.\n"
         "• Its impact spans every industry and aspect of daily life.\n"
         "• While challenges exist, the potential for positive transformation is immense.\n"
         "• Staying informed and adaptable is key to thriving in the AI era.")
    ]

    bullet_slide_layout = prs.slide_layouts[1]
    
    for title_text, content_text in slides_data:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = title_text
        
        # We need to split by lines to add paragraphs
        tf = body_shape.text_frame
        tf.clear() # Clear existing text
        
        lines = content_text.strip().split('\n')
        for i, line in enumerate(lines):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = line.replace('• ', '')
            p.level = 0
            p.font.size = Pt(22)

    # Q&A Slide
    qa_slide = prs.slides.add_slide(prs.slide_layouts[0])
    qa_slide.shapes.title.text = "Questions & Answers"
    qa_slide.placeholders[1].text = "Thank you for your time!"

    prs.save("AI_Impact_Presentation.pptx")
    print("Presentation created successfully as AI_Impact_Presentation.pptx")

if __name__ == '__main__':
    create_presentation()
